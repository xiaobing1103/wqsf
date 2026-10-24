"""生成客户及后台真实截图手册。截图只来自项目页面，不绘制界面。"""
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / 'docs/manual-source/screenshots'
OUT = ROOT / 'docs/deliverables'

def make(title, subtitle):
    d = Document()
    s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.6);s.bottom_margin=Inches(.6)
    s.left_margin=Inches(.65);s.right_margin=Inches(.65)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=d.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
        st.font.underline=False;st.font.italic=False
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
        for b in list(st.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.line_spacing=1.18
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(24)
    d.styles['Heading 1'].font.size=Pt(18)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(0)
    d.styles['Caption'].font.size=Pt(9)
    d.core_properties.title=title;d.core_properties.author='万企商服通'
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run('万企商服通  ').font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    d.add_paragraph(title,'Title');d.add_paragraph(subtitle,'Subtitle')
    d.add_paragraph('版本 V1.1    2026年10月15日')
    return d

def shot(d,name,caption,mobile=False,top=0,bottom=0,left=0,right=0):
    path=SHOTS/(name+'.jpg')
    with Image.open(path) as im: w,h=im.size
    ratio=(h*(1-top-bottom))/(w*(1-left-right))
    width=min(2.65 if mobile else 6.9, (5.5 if mobile else 5.4)/ratio)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
    shape=p.add_run().add_picture(str(path),width=Inches(width),height=Inches(width*ratio))
    shape._inline.docPr.set('descr',caption)
    if any([top,bottom,left,right]):
        bf=shape._inline.xpath('.//pic:blipFill')[0]
        crop=OxmlElement('a:srcRect')
        for k,v in [('t',top),('b',bottom),('l',left),('r',right)]: crop.set(k,str(round(v*100000)))
        bf.insert(1,crop)
    p=d.add_paragraph(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def page(d,title,name,steps,note='',mobile=False,**crop):
    h=d.add_heading(title,1);h.paragraph_format.page_break_before=True
    shot(d,name,title+'  项目实际界面'+(' H5移动端' if mobile else ''),mobile,**crop)
    for i,t in enumerate(steps,1):d.add_paragraph(f'{i}  {t}')
    if note:d.add_paragraph('提醒：'+note)

client=make('万企商服通客户图文操作手册','登录 提交资料 补件与查询进度')
client.add_paragraph('第一次使用，请先向服务人员取得客户账号。按本手册完成登录、资料保存和正式提交；之后可查看进度，按通知要求补充资料。')
shot(client,'client_home','首页  项目实际界面 H5移动端',True)
client.add_paragraph('阅读顺序：登录 → 我的与报单 → 选择服务 → 保存与上传 → 提交确认 → 补件 → 品类参考与企业资料 → 退出切换。')
client.add_paragraph('截图来自当前项目H5移动端实际页面，微信小程序的原生导航和文件选择弹窗可能不同。示例企业为本地演示，按钮与服务以实际开放内容为准。')
page(client,'1 登录客户账号','client_login',[
    '首次使用点“账号登录”，填写服务人员发放的客户用户名和密码，点“登录账号”。不要使用后台管理员账号。',
    '正式微信中可按页面选择绑定当前微信。已经绑定的账号，下次可使用微信一键登录。',
    '没有账号、忘记密码或绑定冲突，联系服务人员。微信昵称不是登录用户名。'],
    '截图中的“本地模拟微信登录”是开发环境入口，不是正式微信授权；正式微信登录需服务方完成配置。',True)
page(client,'2 从我的进入报单列表','client_mine',[
    '点底部“我的”，先确认企业名称正确。点“企业资料”核对信息。',
    '点“全部报单”“处理中”“待补资料”或“已完成”，进入对应列表。',
    '点“消息通知”查看补件与办理进度消息。'],
    '报单列表是二级页。企业或列表不正确时先联系服务人员核对账号关联，不重复建企业。',True)
page(client,'3 找到草稿与已提交报单','client_cases',[
    '按服务、月份、报单编号找到目标，点“查看资料状态”。',
    '草稿显示“尚未提交”和最近保存时间，可进入继续填写。',
    '已提交报单看“最近提交”时间和状态；需要跟进时提供报单编号。'],
    '处理中分类可能包括等待审核等未办结状态，具体以卡片上的状态为准。',True)
page(client,'4 选择服务','client_service',[
    '首页点“全部服务”或目标服务，核对本次服务和资料数量。',
    '点“确认并填写资料”，按页面要求确认企业及月份，进入草稿。',
    '已经有草稿时优先从报单列表恢复，不为了继续填写再创建一笔。'],
    '商标和资质为咨询申请，按表单填写法人姓名、公司名称、联系电话并阅读同意项，可保存草稿或提交咨询；不表示已取得资质。',True)
page(client,'5 按资料清单逐项填写','client_materials',[
    '清单顶部显示报单编号、状态和草稿时间，先核对是否为本次报单。',
    '点对应资料项，填写开票信息、文字或链接，或选择附件。',
    '每项保存后返回清单检查进度。底部“继续补充（N）”表示还缺N项。'],
    '当前贸易增量为11项，以报单实际清单为准；已保存不等于正式提交。',True)
page(client,'6 填写并保存开票信息','client_invoice',[
    '按星号必填项填写，信息应与企业实际资料一致。',
    '填写后点“保存开票信息”或“保存并继续”，等待保存成功反馈。',
    '敏感字段保存后可能只显示掩码，不要因此反复提交。'],
    '不要填写网银、U盾、交易密码或验证码。需要Excel时使用服务人员提供的企业入库表，不混用后台批量导入模板。',True)
page(client,'7 选择附件并确认上传成功','client_upload',[
    '先阅读每项格式、大小和最多数量，再点“选择文件”。',
    '选择正确文件，等待上传完成，确认文件出现在该项附件列表。',
    '已上传图片可点击查看放大；文档按支持的方式打开。失败时按提示重试。'],
    '选择文件不等于上传成功。不同资料项限制不同，微信的文件选择窗口与H5不同。',True)
page(client,'8 正式提交与状态确认','client_submitted',[
    '草稿清单全部必填资料保存后，点“提交审核”，等待成功提示。',
    '进入报单列表再打开该笔，确认“待审核”与最近提交时间。',
    '已提交资料通常只读，可点击资料项查看；100%仅表示资料齐全，不表示审核通过。'],
    '本图是已提交报单的实际只读页面，不是本轮新提交结果；仍是草稿时需检查是否只做了保存。',True)
page(client,'9 补件通知与重新提交','client_notifications',[
    '收到补件提醒后点“查看相关报单”，也可从“我的 → 待补资料”进入。',
    '阅读具体原因，补充对应项，确认上传或保存成功。',
    '返回清单再次“提交审核”，核对状态回到“待审核”，等待复核。'],
    '截图是历史通知，不代表当前待补状态。已完成报单的历史补件通知仍可能保留，以报单当前状态为准。',True)
page(client,'10 经营范围品类参考','client_category',[
    '进入“品类匹配”，可手工填写经营范围；图片识别开放时可上传营业执照。',
    '核对识别文字，有误先校正，再点“开始匹配品类”。',
    '按完整关键词查看推荐结果；OCR不可用时可手工填写。'],
    '推荐只是经营范围关键词参考，不是保险保单识别，也不表示取得经营许可。图片发送前确认授权。',True)
page(client,'11 查看企业资料','client_company',[
    '“我的 → 企业资料”核对公司名称、联系人及所在地区。',
    '信用代码和电话脱敏显示，修改请联系服务人员。',
    '联系客服在微信中且入口开放后使用；H5可能提示在微信中打开。'],
    '反馈问题时提供时间、页面、报单编号、提示及脱敏截图，不提供密码或验证码。',True)
page(client,'12 退出与登录其他账号','client_logout',[
    '切换前先保存资料，在“我的”点“退出登录 / 登录其他账号”。',
    '确认“退出并切换账号”，返回登录页后选“账号登录”使用其他客户账号。',
    '服务器已保存的草稿保留，本机登录及业务缓存清理；未保存输入不保证保留。'],
    '微信一键登录仍会进入原绑定账号。只想返回当前账号时点“取消”。',True)
client.save(OUT/'万企商服通客户图文操作手册.docx')

admin=make('万企商服通后台图文操作手册','客户维护 资料审核 办理与系统管理')
admin.add_paragraph('适用运营、审核人员和超级管理员。每天先找待审核报单，核对资料；合格逐项通过后开始办理，有问题退回指定材料，实际办结后再确认完成。')
shot(admin,'admin_dashboard','工作台  项目后台实际界面')
admin.add_paragraph('阅读顺序：登录 → 客户与账号 → 报单审核与补件 → 线下交接与完成 → 导出 → 模板 → 团队权限 → 功能开关。')
admin.add_paragraph('截图来自当前本地项目，企业示例使用本地演示数据；截图状态不是生产状态。本手册不包含任何登录密码。按钮可用性以角色、资料及报单当前状态为准。')
page(admin,'1 登录后台与每日检查','admin_login',[
    '输入授权的后台用户名及密码，点“登录并进入工作台”。客户账号不能代替管理员账号。',
    '工作台查看待审核、待补资料、处理中和异步任务；进入报单页处理，不只看统计。',
    '左侧进入各模块；右上角“退出登录”用于退出，账号名称不是退出按钮。'],
    '后台通知图标目前未接入完整通知中心。连续密码错误后停止重试，不绕过限流。',top=.2,bottom=.2)
page(admin,'2 客户资料与批量导入','admin_companies',[
    '“客户资料 → 新增客户”用于单个新增。列表可搜索企业名称。',
    '批量操作先“下载导入模板”，保留原列名填写，再点“选择Excel文件”。',
    '等待任务结果并核对成功、失败行；回列表搜索新企业，不把选择文件当作导入完成。'],
    '后台导入模板与客户企业入库信息表用途不同；导入企业不等于已发放登录账号。',bottom=.4)
page(admin,'3 三个点的查看编辑与删除','admin_company_menu',[
    '目标企业行右侧三个点 → “查看详情”，核对企业编号及信息。',
    '“编辑资料”修改后保存；敏感字段不回显时留空表示保留原值。',
    '“删除客户”仅超级管理员使用，先核对未办业务与关联账号，再阅读确认提示。'],
    '删除为软删除，无普通恢复按钮，误删交维护人员处理。不要因客户登录失败删除并重建企业。',bottom=.4)
page(admin,'4 新增或编辑企业资料','admin_company_form',[
    '填写企业名称和联系电话等必填字段，选填企业代码、法人、地区、邮箱和地址。',
    '检查目标企业及输入，再点“保存客户资料”。取消则不保存。',
    '保存后回列表搜索并打开详情核对；提示版本冲突时重新加载最新资料。'],
    '不要在地址或其他字段填密码、验证码和真实认证秘密。')
page(admin,'5 发放客户账号与重置密码','admin_client_accounts',[
    '企业三个点 → “登录账号”，核对弹窗顶部企业名称和已有账号。',
    '创建并关联新账号：填写姓名、用户名、手机号、至少8位初始密码及企业成员身份，点“创建客户账号”。',
    '已有账号忘记密码时点其“重置密码”，确认新密码并通过安全渠道交付客户。'],
    '不把密码写到补件原因或导出表；客户账号不是后台管理员账号，不因“只读成员”文字就承诺所有写入入口都禁用。')
page(admin,'6 找到客户刚提交的报单','admin_cases',[
    '客户提供报单编号后，在“报单管理”搜索，清除旧条件或选择“待审核”。',
    '核对企业、月份、编号、状态与“最近提交”时间，注意分页。',
    '点“查看详情”。草稿未提交不可当作待审核申请；旧报单今天重提仍是原编号。'],
    'TRADE_INCREMENT为贸易增量，IP为商标咨询，QUALIFICATION为资质申报。不要只按创建日期寻找申请。',bottom=.22)
page(admin,'7 资料合格逐项审核通过','admin_review',[
    '左侧“资料审核”逐项读原文及“查看附件”，核对完整性和业务要求。',
    '已提交且合格的资料点其“审核通过”；已提交11/11不等于通过11/11。',
    '所有必填项通过后，右侧“审核通过，开始办理”才满足条件。'],
    '当前图中通过0项，开始办理按钮灰色是正常限制，不应为解锁按钮批准不合格材料。',top=.3,bottom=.13)
page(admin,'8 不合格资料退回补件','admin_review_actions',[
    '勾选问题项旁“需补充”，可选多项，只选择需要客户修改的材料。',
    '右侧“补充原因（客户可见）”写明哪项、哪里不对和如何改。',
    '点“退回所选N项并通知客户”，核对后确认。客户补件重提，再复核对应项。'],
    '未选问题项或未写原因时按钮禁用。原因可写“图片右侧被裁切，请重传四角完整的清晰图片”，不能写密码。',top=.33,bottom=.01,left=.71,right=.04)
page(admin,'9 开始办理 完成与凭证交接','admin_review',[
    '全部必填项通过后点“审核通过，开始办理”，核对状态变“处理中”。',
    '实际业务完成后再点“确认办理完成”，核对“已完成”，不能以资料100%代替办结。',
    '线下凭证交接登记网银、U盾、交易授权凭证状态，再“保存交接状态”。'],
    '仅记待交接、已交接、核验、归还或无需交接等状态，不保存真实密码或验证码。截图为待审核界面，完成按钮仅在对应状态显示。',top=.32,bottom=.34,left=.71,right=.04)
page(admin,'10 导出审核表与资料包','admin_exports',[
    '从报单详情点“导出审核表”，贸易增量可按权限点“下载资料ZIP”。',
    '进入“导出中心”，生成中刷新任务，成功后下载并核对编号。',
    '任务长期未完成或失败，提供任务编号给维护人员；只给授权对象交付资料。'],
    '截图当前无任务。记录按本浏览器会话回读，换浏览器可能不显示旧记录；链接有时效，不传播临时地址。',bottom=.38)
page(admin,'11 核对业务资料模板','admin_template',[
    '“业务模板”点目标服务的“查看模板”，核对模板名称、版本和资料项。',
    '检查必填/选填、敏感标记和数量；当前页面主要查看，非完整编辑发布页面。',
    '需要调整交维护人员按版本规范处理。历史报单使用创建时快照。'],
    '模板已发布不等于客户入口已开放，还受小程序功能开关影响。')
page(admin,'12 团队与权限','admin_team',[
    '超级管理员进入“团队与权限”，点“新增成员”，填写姓名、用户名、手机号、至少12位密码和岗位。',
    '点“X项权限”查看权限清单；其他成员可“调整角色”，按岗位保存。',
    '当前账号不可在这里调整自己的角色。变更后必要时重新登录刷新菜单。'],
    '不要给所有成员超级管理员。当前没有完整成员删除、停用及后台密码重置表单，相关需要交维护人员。',bottom=.3)
page(admin,'13 超级管理员调整功能开关','admin_capabilities',[
    '进入“小程序功能开关”，核对功能与影响范围，只超级管理员可操作。',
    '点目标开关，填调整原因，点“确认保存”并确认二次提示。',
    '核对“已配置开启/关闭”、实际生效及版本；客户重新进页面读取新状态，不需手改数据库或每次重启。'],
    'OCR依赖品类参考且需真实OCR配置，敏感资料仍需相应服务开启，客服需平台配置。关闭会限制草稿及补件修改；开关不是审核通过证明。')
h=admin.add_heading('14 常见问题与交接记录',1);h.paragraph_format.page_break_before=True
for title,txt in [
    ('找不到今天申请','按编号查，清除筛选和检查分页，区分草稿创建与最近提交。客户只保存未提交时不会进入待审核。'),
    ('资料齐全但不能开始办理','左侧每项必填资料必须审核通过，100%是提交完成度，不是审核通过率。也需核对账号权限。'),
    ('退回按钮不可用','至少勾选一个需补充项并填写原因，核对报单当前状态及权限。'),
    ('开关或资料保存冲突','刷新读取最新版本，核对其他人员修改后再操作，不能自动覆盖。'),
    ('附件打开失败','返回详情重新获取预览；仍失败提供报单编号、发生时间和错误提示，交维护人员检查权限与存储。'),
    ('交接记录模板','日期：____  企业与报单编号：____  当前状态：____\n通过项：____  待补项与原因：____  下一步人员与时间：____\n客户通知情况：____  线下交接状态：____（不填密码）'),
    ('每日收尾','跟进待补资料和处理中，检查失败导出，离开共用设备退出登录。按制度保护下载材料，反馈截图脱敏。')]:
    admin.add_heading(title,2);admin.add_paragraph(txt)
admin.add_paragraph('权限不足、需要通用取消报单、分配经办人、完整产品发布或独立审计查询时，请交维护人员处理；当前后台未提供这些完整页面，不在日常操作中手改数据库。')
admin.save(OUT/'万企商服通后台图文操作手册.docx')
print('真实截图客户版与后台版生成完成')
