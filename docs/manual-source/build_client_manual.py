"""旧版示意图构建器已停用，使用真实截图双手册构建器。"""
if __name__ == '__main__':
    import runpy
    from pathlib import Path as _Path
    runpy.run_path(str(_Path(__file__).with_name('build_real_manuals.py')), run_name='__main__')
    raise SystemExit(0)
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / 'docs/manual-source/figures'
OUTPUT = ROOT / 'docs/deliverables/万企商服通客户图文操作手册.docx'
ASSETS.mkdir(parents=True, exist_ok=True)
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
FONT = 'C:/Windows/Fonts/msyh.ttc'
BOLD = 'C:/Windows/Fonts/msyhbd.ttc'
TEAL = '#195456'
GRAY = '#667578'
GOLD = '#EDB84F'

def font(size=24, bold=False):
    return ImageFont.truetype(BOLD if bold else FONT, size)

def text(draw, xy, s, size=24, color='#20383A', bold=False):
    draw.text(xy, s, font=font(size, bold), fill=color)

def rect(draw, box, fill='white', outline='#D6E1E1', radius=15):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)

def badge(draw, xy, n):
    x,y=xy
    draw.ellipse((x,y,x+38,y+38), fill=GOLD)
    draw.text((x+19,y+19),str(n),font=font(23, True),fill='#162E31',anchor='mm')

def checked(draw, x, y, label, size=22):
    draw.rectangle((x,y+5,x+21,y+26),outline=TEAL,width=2)
    draw.line([(x+4,y+15),(x+9,y+21),(x+18,y+9)],fill=TEAL,width=3)
    text(draw,(x+30,y),label,size)

def phone(draw, x, title, tab=False):
    rect(draw,(x,22,x+410,656),fill='#F4F8F7',outline=TEAL,radius=24)
    rect(draw,(x+7,29,x+403,85),fill='white',outline='white',radius=18)
    text(draw,(x+24,42),title,25,bold=True)
    text(draw,(x+338,42),'•••',23)
    if tab:
        rect(draw,(x+8,603,x+402,648),fill='white',outline='white')
        text(draw,(x+42,611),'首页',20)
        text(draw,(x+157,611),'品类匹配',20)
        text(draw,(x+314,611),'我的',20)

def field(draw,x,y,label,value='请按实际资料填写'):
    text(draw,(x,y),label,21,bold=True)
    rect(draw,(x,y+34,x+350,y+85),fill='white',radius=10)
    text(draw,(x+13,y+46),value,20,color=GRAY)

def button(draw,x,y,label,n=None,w=350):
    rect(draw,(x,y,x+w,y+53),fill=TEAL,outline=TEAL,radius=12)
    draw.text((x+w/2,y+26),label,font=font(23,True),fill='white',anchor='mm')
    if n: badge(draw,(x+w-20,y-14),n)

def row(draw,x,y,label,desc='',n=None):
    rect(draw,(x,y,x+350,y+72),radius=12)
    text(draw,(x+16,y+12),label,23,bold=True)
    if desc: text(draw,(x+16,y+43),desc,18,color=GRAY)
    text(draw,(x+326,y+22),'›',25)
    if n: badge(draw,(x+324,y-13),n)

def canvas():
    im=Image.new('RGB',(1000,700),'white')
    return im,ImageDraw.Draw(im)

def save(im,name):
    path=ASSETS/(name+'.png')
    im.save(path)
    return path

def build_figures():
    figs=[]
    im,d=canvas()
    for i,(title,sub) in enumerate([('取得客户账号','由服务人员发放'),('填写并保存资料','草稿可以继续编辑'),('提交审核','确认变为待审核'),('补件或等待办理','有问题按原因重提')]):
        y=45+i*145
        badge(d,(80,y+25),i+1)
        text(d,(145,y+15),title,34,bold=True)
        text(d,(145,y+67),sub,25,color=GRAY)
        if i<3: text(d,(95,y+102),'↓',26,color=TEAL)
    figs.append(save(im,'01_操作路线'))

    im,d=canvas(); phone(d,45,'客户账号登录');phone(d,545,'微信一键登录')
    text(d,(82,108),'账号登录',28,bold=True);badge(d,(396,101),1)
    field(d,75,163,'用户名','管理员为你发放的用户名')
    field(d,75,273,'密码','••••••••')
    checked(d,80,392,'同时绑定当前微信');badge(d,(396,385),2)
    text(d,(80,429),'正式微信中可选择绑定',19,color=GRAY)
    button(d,75,485,'登录账号',3)
    text(d,(580,161),'已绑定账号',30,bold=True)
    text(d,(580,217),'下次可以使用微信登录',24)
    button(d,575,295,'微信授权登录',4)
    row(d,575,387,'首次使用或登录其他账号')
    text(d,(580,505),'未绑定时先取得客户账号',21,color=GRAY)
    text(d,(580,542),'不能用微信昵称作为用户名',21,color=GRAY)
    figs.append(save(im,'02_双登录'))

    im,d=canvas();phone(d,45,'我的',True);phone(d,545,'全部报单')
    row(d,75,111,'示例企业','企业资料',1)
    for i,label in enumerate(['全部报单','处理中','待补资料','已完成']):
        xx=75+(i%2)*180; yy=213+(i//2)*94
        rect(d,(xx,yy,xx+166,yy+76));text(d,(xx+14,yy+13),label,21,bold=True)
        text(d,(xx+14,yy+44),'点击查看',18,color=GRAY)
    badge(d,(389,202),2)
    row(d,75,423,'消息通知','查看补件和进度消息',3)
    row(d,75,510,'退出登录 / 登录其他账号')
    for ix,(yy,st,n) in enumerate([(116,'草稿',4),(240,'待审核',None),(364,'待补资料',None)],1):
        rect(d,(575,yy,925,yy+100));text(d,(591,yy+12),'示例企业 · 贸易增量',22,bold=True)
        text(d,(591,yy+46),f'报单编号 DEMO00{ix}',19,color=GRAY)
        text(d,(591,yy+74),st,20,color=TEAL)
        if n:badge(d,(897,yy-10),n)
    text(d,(580,520),'先核对编号，再进入报单',23,color=GRAY)
    figs.append(save(im,'03_列表入口'))

    im,d=canvas();phone(d,45,'服务选择');phone(d,545,'资料提交清单')
    text(d,(78,117),'先确认企业和月份',25,bold=True)
    row(d,75,175,'贸易增量','开放后可申请',1)
    row(d,75,269,'商标注册咨询')
    row(d,75,363,'资质申报服务')
    text(d,(80,502),'没有服务时请联系服务人员',20,color=GRAY)
    text(d,(580,110),'按当前清单准备材料',25,bold=True)
    row(d,575,170,'开票资料','点击填写',2)
    row(d,575,269,'文件资料','格式数量以页面要求为准',3)
    row(d,575,368,'文字或链接资料','填写后执行保存')
    button(d,575,517,'继续补充',4)
    figs.append(save(im,'04_资料清单'))

    im,d=canvas();phone(d,45,'开票信息');phone(d,545,'附件资料')
    field(d,75,119,'公司名称','与实际资料保持一致')
    field(d,75,223,'税号与联系信息','按页面必填项逐项填写')
    field(d,75,327,'开户行与基本户账号','保存后可能脱敏显示')
    button(d,75,482,'保存开票信息',1)
    text(d,(80,563),'不要录入密码或验证码',21,color='#98572A')
    text(d,(580,118),'先看格式 大小 数量要求',23,bold=True)
    button(d,575,198,'选择附件',2)
    row(d,575,304,'示例资料.png','上传成功后才会显示',3)
    text(d,(580,412),'点击文件查看图片大图',24,bold=True);badge(d,(892,403),4)
    text(d,(580,470),'未上传成功时请重试',21,color=GRAY)
    text(d,(580,515),'已提交资料通常只读',21,color=GRAY)
    figs.append(save(im,'05_保存与附件'))

    im,d=canvas();phone(d,45,'资料提交清单');phone(d,545,'报单状态')
    text(d,(80,121),'资料已准备齐全',28,bold=True)
    text(d,(80,180),'先检查必填项和上传结果',22,color=GRAY)
    row(d,75,247,'已保存资料说明','保存不等于正式提交',1)
    button(d,75,378,'提交审核',2)
    text(d,(80,491),'缺少资料时先继续补充',22,color=GRAY)
    text(d,(580,118),'提交后再次查看报单',26,bold=True)
    row(d,575,192,'报单编号 DEMO001','状态 待审核',3)
    row(d,575,304,'最近提交时间','核对是否是刚才提交的时间',4)
    text(d,(580,426),'待审核 = 等待后台审核',23,bold=True)
    text(d,(580,482),'不是草稿，也不是已完成',22,color=GRAY)
    figs.append(save(im,'06_提交确认'))

    im,d=canvas();phone(d,45,'消息通知');phone(d,545,'待补资料')
    row(d,75,129,'资料需要补充','点击进入相关报单',1)
    text(d,(80,252),'示例补充原因',25,bold=True)
    text(d,(80,311),'图片右侧被裁切',23)
    text(d,(80,353),'请上传完整清晰的图片',23)
    text(d,(80,490),'也可从 我的 → 待补资料进入',20,color=GRAY)
    row(d,575,129,'问题资料项','查看具体补充说明',2)
    button(d,575,258,'选择附件',3)
    row(d,575,363,'新的资料图片','确认上传并保存成功')
    button(d,575,508,'提交审核',4)
    figs.append(save(im,'07_补件重提'))

    im,d=canvas();phone(d,45,'咨询申请');phone(d,545,'品类匹配')
    field(d,75,114,'法人姓名','按实际信息填写')
    field(d,75,211,'公司名称','请输入公司全称')
    field(d,75,308,'联系电话','请输入联系电话')
    checked(d,80,415,'阅读并同意页面说明',21)
    button(d,75,478,'提交咨询',1)
    field(d,575,113,'经营范围','软件开发；电子产品销售')
    text(d,(580,232),'可手工填写或校正识别文字',20,color=GRAY)
    text(d,(580,302),'有图片识别入口时',23,bold=True)
    text(d,(580,349),'先上传，再核对识别结果',21)
    text(d,(580,410),'提交匹配后查看推荐结果',23,bold=True);badge(d,(895,402),2)
    text(d,(580,489),'结果仅供经营范围参考',21,color=GRAY)
    text(d,(580,533),'不表示取得经营许可',21,color=GRAY)
    figs.append(save(im,'08_咨询与参考'))

    im,d=canvas();phone(d,45,'我的');phone(d,545,'切换账号')
    row(d,75,120,'企业资料','需要修改请联系服务人员',1)
    row(d,75,221,'消息通知')
    row(d,75,322,'联系客服','微信中开放后可使用',2)
    row(d,75,423,'退出登录 / 登录其他账号','先保存正在填写的资料',3)
    text(d,(580,149),'退出当前账号',28,bold=True)
    text(d,(580,218),'服务器上已保存的草稿保留',20)
    text(d,(580,260),'本机登录及业务缓存会清理',20)
    button(d,575,359,'退出并切换账号',4)
    text(d,(580,485),'改用其他账号时选择账号登录',20,color=GRAY)
    figs.append(save(im,'09_企业与退出'))
    return figs

doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.65);s.bottom_margin=Inches(.65)
s.left_margin=Inches(.75);s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    st.font.underline=False
doc.styles['Subtitle'].font.italic=False
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.3
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('万企商服通客户操作手册  ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='万企商服通客户图文操作手册'
doc.core_properties.author='万企商服通'

def para(t,bold=False):
    p=doc.add_paragraph();r=p.add_run(t);r.bold=bold;return p

def picture(path,caption,width=6.6):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(0)
    r=p.add_run();shape=r.add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    p=doc.add_paragraph(caption+'  操作示意',style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def page(title,img,intro,steps,tips):
    h=doc.add_heading(title,level=1);h.paragraph_format.page_break_before=True
    para(intro);picture(img,title)
    for i,t in enumerate(steps,1): para(f'{i}  {t}')
    for t in tips: para(t)

figs=build_figures()
doc.add_paragraph('万企商服通客户图文操作手册',style='Title')
doc.add_paragraph('企业资料提交与办理进度查看',style='Subtitle')
para('版本 V1.0    2026年10月14日')
para('本手册帮助你完成客户登录、填写资料、提交审核、补充材料和查看结果。首次使用请先向服务人员取得客户账号；不要使用后台管理员账号。')
picture(figs[0],'从取得账号到提交审核的操作路线',5.4)
para('按需求翻页',True)
para('第2页 登录与绑定微信   第3页 找到企业和报单   第4页 准备资料')
para('第5页 保存信息与上传附件   第6页 提交并确认   第7页 补件重提')
para('第8页 咨询与品类参考   第9页 客服与切换账号   第10页 状态与常见问题')
para('图中编号对应本页操作步骤。页面名称和按钮以实际开放内容为准，配图标有“操作示意”，不要照抄图中示例企业或报单编号。')

page('1 登录账号与绑定微信',figs[1],
     '首次使用先选账号登录。已经绑定的客户，下一次可以使用微信一键登录。',
     ['选择“账号登录”，填写服务人员发放的用户名和密码。',
      '正式微信中可勾选“同时绑定当前微信”。只想登录其他账号时，请谨慎选择绑定。',
      '点击“登录账号”，成功后进入首页。',
      '已绑定客户可点击“微信授权登录”；未绑定时先完成账号登录和绑定。'],
     ['没有账号或忘记密码，请联系服务人员。微信昵称不是客户用户名。',
      '看到“本地模拟微信登录”表示开发测试模式，不是正式微信授权。绑定冲突时不要反复重试，请联系服务人员核对。'])

page('2 找到企业和自己的报单',figs[2],
     '“我的”是查看企业、报单和通知的入口。先确认企业正确，再打开具体报单。',
     ['进入“我的”，点击企业卡片或“企业资料”查看当前账号的企业信息。',
      '点击“全部报单”“处理中”“待补资料”或“已完成”，进入对应报单列表。',
      '点击“消息通知”查看补件或进度消息。',
      '在列表点击目标报单，核对编号、服务类型、月份和状态。'],
     ['同一家企业可以有多个报单。联系服务人员时请提供报单编号，不要只说企业名称。',
      '报单列表是二级页面，不在底部导航中。企业信息不正确或列表为空时，先让服务人员核对账号关联。'])

page('3 选择服务并准备资料',figs[3],
     '需要新申请时，从首页进入服务选择。已有草稿请先从报单列表恢复，避免重复创建。',
     ['选择需要的服务，按页面确认企业和报单月份，创建或恢复草稿。',
      '进入“资料提交清单”，点开开票资料并逐项填写。',
      '按清单打开文件、文字或链接资料，阅读具体要求后填写或上传。',
      '返回清单检查缺失项；提示“继续补充”时先补齐，不要以为已经提交。'],
     ['资料名称、数量、必填项以当前报单清单为准，图中只展示常见类型。',
      '服务入口未出现或显示暂未开放，请联系服务人员。不要新建另一种服务来绕过限制。'])

page('4 保存信息与上传附件',figs[4],
     '选择文件不等于上传成功，填完表格也不等于保存成功。每次操作后都检查结果。',
     ['填写开票信息，点击对应保存按钮，等待“已保存”等成功反馈。',
      '先看允许格式、大小和数量，再点击附件选择按钮。',
      '选择正确文件，等待上传完成；文件应出现在附件列表。',
      '点击已上传文件检查内容。图片可预览放大；文档按支持的打开方式查看。'],
     ['需要填写Excel时，使用服务人员提供的企业入库表，填写后按指定资料项上传；当前没有统一的所有模板下载入口。',
      '不要提交网银密码、U盾密码、交易密码或验证码。敏感字段保存后可能脱敏显示，不表示保存失败。'])

page('5 提交审核并确认成功',figs[5],
     '保存资料只是准备草稿。点击“提交审核”并看到待审核状态，才完成本次提交。',
     ['逐项核对资料已保存、必填项已补齐。草稿时间不等于正式提交时间。',
      '点击“提交审核”，等待成功提示，不要连续重复点击。',
      '进入“我的 → 全部报单”，打开该报单，确认状态为“待审核”。',
      '核对报单编号和最近提交时间。需要跟进时将编号告知服务人员。'],
     ['仍显示“草稿”时，先确认是否只保存了资料。出现缺失提示时返回清单补充。',
      '资料完成度100%只表示提交齐全，不表示已审核通过或业务已办理完成。已提交资料通常只读，等待后台审核。'])

page('6 按要求补充资料并重新提交',figs[6],
     '看到“待补资料”表示需要处理具体问题。先读原因，再修改对应材料。',
     ['从“消息通知”或“我的 → 待补资料”打开报单。',
      '阅读补充原因，确认哪一项有问题以及要求怎样修改。',
      '替换或补充对应附件、文字等资料，检查保存和上传结果。',
      '返回清单再次点击“提交审核”，确认状态回到“待审核”。'],
     ['其他已通过的资料无需无故重复上传。提示图片模糊时，请重新提供完整、清晰的图片。',
      '报单只读或按钮不可用时，可能是服务关闭或状态已变化。联系服务人员处理，不重复建单。'])

page('7 咨询申请与品类参考',figs[7],
     '商标、资质咨询与品类参考是不同功能。只操作与你的实际需求一致的服务。',
     ['咨询申请：填写法人姓名、公司名称和联系电话，阅读同意项后“提交咨询”；未准备好可“保存草稿”。',
      '品类参考：输入经营范围并提交匹配。有图片识别入口时，上传后先核对识别文字，再查看推荐结果。'],
     ['咨询提交后，到报单列表查看编号和状态。咨询申请不表示已经注册商标或取得资质。',
      '图片识别无法使用时可手工填写经营范围。品类推荐仅是关键词参考，不是经营许可或保险保单识别。无匹配结果时检查经营范围用词。'])

page('8 企业资料 客服与切换账号',figs[8],
     '企业资料可查看，修改和账号问题请联系服务人员。离开共用设备前请退出登录。',
     ['进入“我的 → 企业资料”，核对企业名称及联系信息；有误联系服务人员更正。',
      '客服入口开放时，在微信小程序点“联系客服”。H5提示在微信中打开。',
      '切换账号前先保存资料，再点“退出登录 / 登录其他账号”。',
      '确认“退出并切换账号”，返回登录页后使用另一个客户账号登录。'],
     ['服务器上已保存的草稿不会因退出而删除；本机登录、企业选择和业务缓存会清理。未保存输入不保证保留。',
      '微信一键登录仍进入原绑定账号。想登录其他账号请选择“账号登录”。微信客服未配置或入口未开放时，使用原业务联系渠道。'])

h=doc.add_heading('9 状态说明与常见问题',1);h.paragraph_format.page_break_before=True
para('先根据报单状态判断下一步。客户收到补件要求后重新提交，后台会复核材料，再办理业务。')
table=doc.add_table(rows=1,cols=3);table.autofit=False
widths=[1.0,3.6,2.0]
for c,w,t in zip(table.rows[0].cells,widths,['状态','说明','你需要做什么']):
    c.width=Inches(w);c.text=t
states=[('草稿','尚未正式提交','继续填写并提交'),('待审核','已提交 等待后台审核','等待或查看通知'),('待补资料','后台要求补充具体资料','按原因补件重提'),('处理中','进入实际办理阶段','查看进度'),('已完成','本报单已结束办理','查看结果'),('已取消','该报单已终止','联系服务人员核对')]
for vals in states:
    cells=table.add_row().cells
    for c,w,v in zip(cells,widths,vals):c.width=Inches(w);c.text=v
for ri,row in enumerate(table.rows):
    trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
    if ri==0:trPr.append(OxmlElement('w:tblHeader'))
    for ci,c in enumerate(row.cells):
        pr=c._tc.get_or_add_tcPr()
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8ED' if ri==0 else ('F6F8F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
        pr.append(borders)
        margins=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            m=OxmlElement('w:'+side);m.set(qn('w:w'),'95');m.set(qn('w:type'),'dxa');margins.append(m)
        pr.append(margins)
        va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');pr.append(va)
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.2
            if ci==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(10.5);r.bold=ri==0
para('')
doc.add_heading('遇到问题先这样处理',2)
for t in ['忘记密码或没有账号：联系服务人员发放或重置，不反复猜密码。',
          '刚提交但后台找不到：提供报单编号，确认不是仅保存草稿；让对方检查筛选和最近提交时间。',
          '不能上传或修改：检查是否已提交、服务未开放或状态已变化；不要重复创建新单。',
          '附件打不开：返回报单重新点击。仍失败时提供提示截图，不传播临时预览链接。',
          '保存提示数据已变化：返回重新加载，核对最新内容后再操作。']:
    para(t)
doc.add_heading('向服务人员反馈时',2)
para('说明发生时间、所在页面、报单编号、点击步骤、错误提示，并提供脱敏截图。不要提供密码、验证码、登录Token或个人证件原文。')
para('后台流程为逐项审核 → 必填资料通过后开始办理 → 实际办结后确认完成。补件过程中由客户重提后再审核，不需要客户自己切换办理状态。')
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUTPUT)
print(f'输出 {OUTPUT}，嵌入 {len(figs)} 张操作示意图')
